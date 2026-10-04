"""CV reader tool for PDF, DOCX and TXT files (Week 2, Lesson 1: tool use)."""

from pathlib import Path

import pdfplumber
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.tools import tool


MIN_CV_CHARS = 40


def _read_pdf(file_path: Path) -> str:
    """Extract text from a PDF file."""

    pages = []

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)

    return "\n".join(pages)


def _table_lines(table: Table) -> list[str]:
    """One line per table row; merged cells are read once."""

    lines = []

    for row in table.rows:
        cells = []
        seen = set()

        for cell in row.cells:
            if cell._tc in seen:
                continue

            seen.add(cell._tc)

            text = " ".join(
                paragraph.text.strip()
                for paragraph in cell.paragraphs
                if paragraph.text.strip()
            )

            if text:
                cells.append(text)

        if cells:
            lines.append(" | ".join(cells))

    return lines


def _read_docx(file_path: Path) -> str:
    """Extract text from a DOCX file."""

    document = Document(file_path)
    lines: list[str] = []

    for section in document.sections:
        for header in (section.header, section.first_page_header):
            for paragraph in header.paragraphs:
                text = paragraph.text.strip()

                if text and text not in lines:
                    lines.append(text)

            for table in header.tables:
                lines.extend(line for line in _table_lines(table) if line not in lines)

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]

        if tag == "p":
            text = Paragraph(child, document).text.strip()

            if text:
                lines.append(text)

        elif tag == "tbl":
            lines.extend(_table_lines(Table(child, document)))

    return "\n".join(lines)


def _read_txt(file_path: Path) -> str:
    """Read text from a plain text file."""

    return file_path.read_text(encoding="utf-8", errors="replace")


@tool(description="Read a candidate CV from a PDF, DOCX, or TXT file and return its text.")
def read_cv(file_path: str) -> str:
    path = Path(file_path)

    if not path.exists():
        return f"Error: CV file was not found at '{file_path}'."

    if not path.is_file():
        return f"Error: '{file_path}' is not a file."

    extension = path.suffix.lower()

    try:
        if extension == ".pdf":
            text = _read_pdf(path)

        elif extension == ".docx":
            text = _read_docx(path)

        elif extension == ".txt":
            text = _read_txt(path)

        else:
            return (
                "Error: Unsupported CV format. "
                "Please use PDF, DOCX, or TXT."
            )

    except Exception as error:
        return f"Error while reading the CV: {error}"

    text = text.strip()

    if not text:
        return "Error: No readable text was found in the CV."

    if len(text) < MIN_CV_CHARS:
        return (
            f"Error: Only {len(text)} characters could be read from the CV. "
            "If it is a scanned image or built from text boxes, export it as "
            "a text-based PDF or DOCX and try again."
        )

    return text
