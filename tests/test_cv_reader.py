"""Tests for the CV reader tool."""

from docx import Document

from tools.cv_reader import read_cv


def test_docx_header_and_tables_are_read(tmp_path):
    """Many templates keep the name in the header and jobs in tables."""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Sara Nassar | sara@example.com"
    doc.add_paragraph("EXPERIENCE")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "2021-2023"
    table.rows[0].cells[1].text = "Ward Nurse, City Hospital"
    merged = table.rows[1].cells[0].merge(table.rows[1].cells[1])
    merged.text = "Supervised two junior nurses"
    path = tmp_path / "cv.docx"
    doc.save(path)

    text = read_cv.invoke({"file_path": str(path)})

    assert text.splitlines()[0] == "Sara Nassar | sara@example.com"
    assert "2021-2023 | Ward Nurse, City Hospital" in text
    assert text.count("Supervised two junior nurses") == 1


def test_nearly_empty_cv_gets_a_helpful_error(tmp_path):
    path = tmp_path / "scan.txt"
    path.write_text("Name")
    message = read_cv.invoke({"file_path": str(path)})
    assert message.startswith("Error: Only 4 characters")


def test_txt_with_bad_bytes_is_still_read(tmp_path):
    path = tmp_path / "cv.txt"
    path.write_bytes("Jos\xe9 Garc\xeda, enfermero con experiencia hospitalaria y atenci\xf3n al paciente".encode("latin-1"))
    assert "enfermero" in read_cv.invoke({"file_path": str(path)})
